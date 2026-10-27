"""Build the editable author-review report from frozen recorded results."""
import hashlib,json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
def read(name):return json.loads((ROOT/name).read_text())

def main():
    analysis=read('experiments/optimization85/analysis/validation_analysis.json')
    score=read('experiments/optimization85/test-score.json');resource=read('experiments/optimization85/test-resources.json')
    cfg=read('code/checkpoints/candidate85/ensemble.json')
    fingerprint=hashlib.sha256((ROOT/'experiments/optimization85/freeze.json').read_bytes()).hexdigest()
    d=Document();section=d.sections[0]
    section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.78)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
    normal=d.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(6)
    d.styles['Title'].font.size=Pt(23);d.styles['Title'].paragraph_format.space_after=Pt(8)
    d.styles['Subtitle'].font.size=Pt(11)
    for name,size in [('Heading 1',14),('Heading 2',11.5)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(5)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run();r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    d.core_properties.author='Eillot1031';d.core_properties.title='Selecting Reliable LLM Reasoning'
    d.core_properties.subject='DASE7506 Mini Project 2 experimental report'
    def p(text):return d.add_paragraph(text)
    def h(text):d.add_heading(text,1)
    def sub(text):d.add_heading(text,2)
    def page(title):d.add_page_break();h(title)
    def table(headers,rows,widths=None):
        t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        if widths:
            for c,w in zip(t.columns,widths):c.width=Inches(w)
        for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
        for row in rows:
            cells=t.add_row().cells
            for c,text in zip(cells,row):c.text=str(text)
        for i,row in enumerate(t.rows):
            for j,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if widths:cell.width=Inches(widths[j])
                pr=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'244A70' if i==0 else 'F1F4F7' if i%2 else 'FFFFFF');pr.append(shade)
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
                pr.append(borders);margins=OxmlElement('w:tcMar')
                for edge in ('top','bottom','left','right'):
                    node=OxmlElement('w:'+edge);node.set(qn('w:w'),'80' if edge in ('top','bottom') else '95');node.set(qn('w:type'),'dxa');margins.append(node)
                pr.append(margins)
                for par in cell.paragraphs:
                    par.paragraph_format.space_after=Pt(0);par.paragraph_format.space_before=Pt(0);par.paragraph_format.line_spacing=1.05
                    par.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in par.runs:
                        run.font.name='Times New Roman';run.font.size=Pt(9.5);run.font.bold=i==0
                        run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        d.add_paragraph().paragraph_format.space_after=Pt(1)
    def code(text):
        par=d.add_paragraph();par.paragraph_format.line_spacing=1.;par.paragraph_format.space_after=Pt(6)
        run=par.add_run(text);run.font.name='DejaVu Sans Mono';run.font.size=Pt(8.5)
    def figure(filename,caption,width=6.5):
        d.add_picture(str(ROOT/'report/figures85'/filename),width=Inches(width))
        par=p(caption)
        for r in par.runs:r.font.size=Pt(9)

    d.add_paragraph('Selecting Reliable LLM Reasoning','Title')
    d.add_paragraph('DASE7506 Advanced Machine Learning Methods\nMini Project 2    Eillot1031    27 October 2026','Subtitle')
    h('Abstract')
    p(f'A compact ensemble selects a reference-correct numerical answer for {score["correct"]} of {score["n"]} questions in the complete public test bank, giving {score["accuracy"]:.3f}% accuracy. Three candidate verifiers are combined with three seeds of a DeepSets comparison network and a numerical equality check. The full CPU prediction and scoring processes take {resource["total_seconds"]:.3f} seconds locally, with {resource["asset_file_bytes"]/1024**2:.3f} MiB of inference assets. This improves the previous frozen result by two questions; the subsequent 85% aspiration remains two correct questions away.')
    p('Agreement alone is insufficient: difficult-validation accuracy improves from 69.014% for majority voting to 88.451%. The latest extension improves validation by only one question, so evidence for its incremental benefit is limited. High answer-score margins identify an easier subset, but scores are not calibrated probabilities. Seven distinct frozen methods were tested during development; the final result is an adaptive benchmark result rather than an untouched holdout estimate.')
    h('Task and research question')
    p('Each question supplies four generated solution texts and a 1024-dimensional representation for every candidate. The output is one candidate index. The unchanged official parser and scorer compare the selected final numeric answer with the reference; they do not grade every intermediate reasoning step [1, 2]. The model performs offline selection and never calls an LLM during evaluation.')
    p('The research question is whether comparing candidate evidence and checking numerical consistency improves selection beyond consensus, while remaining small and fast enough for the course limits. A related question is whether the resulting scores meaningfully identify uncertainty.')
    table(['Split','Questions','Difficult groups'],[['Training','4,755','2,114'],['Validation','798','355'],['Test','467','467 by protocol']],[2.5,1.6,2.6])
    p('A difficult group has at least two distinct valid parsed answers. Each supplied group has at least one reference-correct candidate. Easier groups explain much of the difference between complete-validation and difficult-validation accuracy.')

    page('Method and implementation')
    sub('Candidate evidence')
    p('The supplied embedding is normalized to length 32. Candidate features concatenate this representation, its difference from the within-question mean, and handcrafted descriptors. The descriptors include valid-answer agreement, logarithmic trace length and answer magnitude, numerical overlap with the question, formatting counts, and generic lexical indicators. Each descriptor is accompanied by its deviation from the group mean. These are general functions of the provided inputs; question IDs are never features.')
    p('Two networks additionally receive the element-wise product of the candidate embedding and the group-mean embedding. This supplies a simple interaction with the surrounding candidates. All features use training means and scales, with a standard-deviation floor of 0.15 and standardized values clipped to [-10, 10]. The final networks have two hidden layers, with widths shown below, ReLU activations and training dropout.')
    table(['Member','Hidden widths','Objective','Seed','Frozen epoch'],[
        ['ema42','128 and 32','BCE with EMA','42','46'],
        ['ema_context73','64 and 32','BCE with EMA','73','8'],
        ['context17','64 and 32','List loss plus BCE','17','132'],
        ['DeepSets x3','64 and 32','BCE with EMA','17, 42, 73','19']],[1.35,1.2,1.65,1.0,1.5])
    p('Training uses AdamW with learning rate 0.001 and dropout 0.4. Weight decay is 0.01 for BCE and 0.1 for list loss; EMA is 0.99 where used. Gradients use the 2,114 difficult training groups and exclude unparseable candidates. The original three models standardize using all training inputs; DeepSets uses difficult training inputs. Following set-based shared encoding and symmetric aggregation [5], its 64-dimensional encoder is compared with the valid-candidate mean, maximum and same-answer mean. Concatenating these with the candidate, its mean difference and product produces 384 features for the 32-to-1 head, preserving candidate permutation equivariance.')
    sub('Answer comparison and arithmetic evidence')
    p('Scores are divided by frozen within-question scales. The original ensemble supplies 90% of the neural score; the mean of three DeepSets seeds supplies 10%. Same-answer evidence then supplies half the candidate score. The revised arithmetic penalty is also averaged within each numeric answer group.')
    code('s_i = 0.9 mean_old(f_i/T) + 0.1 mean_set(f_i/T)\ng_i = 0.5 s_i + 0.5 mean_same(s_j) - 0.5 mean_same(b_j)')
    p('Here b counts numeric equalities with relative error above 0.001, using max(1, |left|, |right|) as denominator. The checker handles fractions, percentages and bounded powers, and avoids treating numeric tails of symbolic terms as full equations. Local-variable evidence was explored but has zero final weight. Unparseable answers are masked. Locally consistent arithmetic does not establish semantic correctness.')

    page('Experimental protocol and test history')
    p('Exploration compared answer-only features, handcrafted descriptors, supplied embeddings, relative embeddings, neural objectives, regularization, random seeds, boosted trees, sparse lexical models and kernel verifiers. Complete configurations and learning curves for 67 completed exploratory runs are retained. Validation selected base models and ensemble combinations, with difficult-validation accuracy as the primary criterion. Forward selection considered at most six members and a short answer-pooling grid.')
    p('GPU five-fold studies keep all four candidates together. The initial 24-family search selected common epochs using five-epoch-smoothed OOF curves; its refits did not exceed the original validation-selected ensemble. A further six DeepSets and six answer-balanced/SAM [6] configurations used nine-epoch smoothing over 100 epochs. Their best smoothed OOF scores were 82.103% and 82.093%. A fixed 17-setting validation blend grid selected the 10% DeepSets contribution. The train-OOF arithmetic proxy improved from 83.964% to 84.248% with the revised checker; this is not an OOF estimate of the final six-network ensemble.')
    table(['Evaluation','Method stage','Test correct','Test accuracy'],[
        ['1','Initial ensemble','383','82.013%'],['2','Hard groups and averaging','389','83.298%'],
        ['3','Single CV family','384','82.227%'],['4','Multiple CV families','385','82.441%'],
        ['5','OOF proxy and arithmetic','387','82.869%'],['6','Original frozen ensemble','393','84.154%'],
        ['7','DeepSets blend and revised checker','395','84.582%']],[.8,3.0,1.2,1.7])
    sub('Interpretation of the development history')
    p('Aggregate test scores monitored the requested 84% target and subsequent 85% aspiration; development continued after earlier scores. This is adaptive test feedback and does not satisfy a strictly one-shot, final-only protocol. No individual test references or error cases were inspected for fitting or feature design. The latest round froze its train/validation-selected candidate before one complete test evaluation and made no subsequent score-based changes. These safeguards do not make the test bank untouched.')
    p('All seven official results and freeze records are retained. The instructor determines eligibility under the strict course rule. An independent audit bank is needed for a clean assessment. Validation also reflects model and blend selection; its one-question latest gain is development evidence, not an independent significance result.')

    page('Ablations and comparative results')
    rows=[[r['name'],f'{r["accuracy"]:.3f}%',f'{r["hard_accuracy"]:.3f}%'] for r in analysis['comparisons']]
    table(['Method','Complete validation','Difficult validation'],rows,[3.3,1.7,1.7])
    figure('validation_comparison.png','Figure 1  Difficult-validation comparisons. Ablations are development runs; the final model also benefits from seed and ensemble selection.',6.4)
    p('Answer agreement provides a substantial baseline. Handcrafted descriptors and embeddings each add predictive information, while the combined model is stronger than either single information source in these runs. This supports using complementary evidence, but it does not establish that the model has learned a faithful mathematical proof checker: formatting, answer statistics and generator-specific patterns remain available shortcuts.')
    aa=analysis['arithmetic_ablation']
    p(f'On the same six frozen networks, removing the arithmetic penalty gives {aa["without"]["hard_accuracy"]:.3f}% difficult-validation accuracy; adding it gives {aa["with"]["hard_accuracy"]:.3f}%, a gain of three validation questions. The entire extension improves the previous complete method by one validation question and two test questions. The answer-balanced/SAM family was rejected: its standalone difficult-validation accuracy was 85.634%, and its tested blends did not exceed the selected method. Adding complexity alone was insufficient.')

    page('Uncertainty and selective accuracy')
    figure('reliability.png','Figure 2  Validation reliability diagnostics. Left: retaining questions with larger distinct-answer margins. Right: uncalibrated softmax confidence over distinct answers.',6.4)
    table(['Coverage on difficult validation','Retained questions','Accuracy'],[
        [f'{r["coverage"]*100:.1f}%',r['n'],f'{r["accuracy"]:.3f}%'] for r in analysis['coverage']],[3.2,1.7,1.8])
    p('The margin is the difference between the best and second-best distinct-answer score. It ranks uncertainty more usefully than treating every selected answer as equally reliable: the highest-margin quarter contains 89 questions and is entirely reference-correct in this validation sample. This is an exploratory observation on the development split; it does not guarantee 100% precision on unseen questions.')
    p(f'The reliability plot uses five equal-count bins and an uncalibrated softmax over distinct-answer scores. Its expected calibration error is {analysis["uncalibrated_ece"]*100:.2f} percentage points. Score scales were normalized for fusion, not fitted as probabilities. These values therefore should not be interpreted as calibrated likelihoods of correctness. A new calibration split or nested calibration procedure would be required before deploying a confidence threshold.')
    diff=analysis['vs_majority_hard']
    p(f'Compared with majority voting, the final method fixes {diff["fixed"]} difficult-validation questions while regressing on {diff["regressed"]}. The net gain is substantial, but individual high-margin errors remain. Correlated candidate traces can agree on the same omitted condition, and a confident learned score can reinforce that shared mistake.')

    page('Failure analysis and limits of correctness labels')
    sub('An omitted quantity despite correct arithmetic')
    p('In a validation question, a class has 250 students, all 140 girls attend, and attending girls are twice as numerous as attending boys. The reference answer is 40 absent boys: 250 - 140 - 70. Three candidates instead return 180 by subtracting only 70 attending boys from the total. The selected wrong trace contains locally correct calculations. The numerical checker cannot detect the missing group of 140 girls; the failure is about representing the question conditions.')
    sub('A plausible reading can disagree with the reference')
    p('Another question says Yun starts with 20 paperclips and loses 12, while Marion has "1/4 more than" Yun now has, plus 7. The supplied reference is 9, following (20 - 12)/4 + 7. The model selects 17, following 8 + 8/4 + 7, which is a natural reading of "one quarter more than". The official scorer correctly marks it wrong under its reference, but the example cautions against equating every reference mismatch with demonstrably invalid reasoning. No reference or scorer was changed.')
    sub('Reversing a relation between entities')
    p('In a pens-and-pencils question with 108 utensils, the reference relation gives 16 pens, while the selected candidate reverses which entity equals five times the other plus 12 and returns 92. Both alternatives contain consistent algebra once the relation is chosen. Distinguishing them requires grounding the equation in the wording, beyond checking arithmetic or consensus.')
    sub('What the method can and cannot establish')
    p('The method is an empirical verifier of final answers within the supplied candidate distribution. It checks only a subset of explicit calculations and cannot prove completeness of a solution, validate arbitrary symbolic algebra, or guarantee the intended interpretation of ambiguous language. Conversely, a correct final answer may contain incorrect intermediate steps. The published metric should be interpreted according to this scope.')
    p('The strongest limitations are adaptive public-test feedback, a relatively small difficult-validation split, generator-specific features, and correlated mistakes across candidates. Increasing network capacity, adding many arithmetic-relation features, or learning a nonlinear score fusion did not reliably improve the train OOF results. Those failures argue for better evidence and cleaner evaluation, rather than further uncontrolled search.')

    page('Resources and reproducibility')
    table(['Resource','Measured value or configuration'],[
        ['Frozen asset files',f'{resource["asset_file_bytes"]:,} bytes  {resource["asset_file_bytes"]/1024**2:.3f} MiB'],
        ['Fresh prediction process',f'{resource["prediction_process_seconds"]:.3f} seconds'],
        ['Independent scoring process',f'{resource["scoring_process_seconds"]:.3f} seconds'],
        ['Total complete-bank CPU evaluation',f'{resource["total_seconds"]:.3f} seconds'],
        ['Evaluation hardware','Ryzen 7 9700X  four library threads'],
        ['Training search GPU','RTX 5060 Ti 16 GB  PyTorch 2.9.0 CUDA 13.0'],
        ['Frozen training reproduction','Original CPU 24.987 s  new GPU seeds 5.061 s']],[3.0,3.7])
    p('These totals include new-process startup, imports, supplied input loading, feature extraction, all ensemble members and independent scoring. They are local measurements; the Xeon Platinum 8457C reference machine remains authoritative for eligibility. The final assets are NumPy arrays stored without compression, including all learned means, scales, weights and score normalization constants.')
    p('The original 67 exploratory runs record 1,010.390 seconds of summed per-run work, not wall-clock duration. The initial GPU search contains 24 families times five folds times 80 epochs and reused caches. The latest two searches each contain six configurations times five folds times 100 epochs; measured times are retained in experiments/optimization85/. A single wall-clock total for all development was not captured. Frozen training reproduction does not repeat model selection.')
    code('python -m pip install -r code/requirements-inference.txt\npython code/evaluate.py --split test \\\n  --checkpoint code/checkpoints/candidate85 \\\n  --output-dir runs/reproduction85')
    p('Use Python 3.12 and NumPy 2.3.5. For training, code/reproduce_training.py rebuilds the original CPU models; code/reproduce_optimization85.py refits the new CUDA seeds at fixed epoch 19. Both reproduced all learned arrays exactly locally. The latter can consume the rebuilt baseline via --base. The supplied parser, data and scorer remain unchanged. Export parity, independent IDs/batches, two candidate permutations and arithmetic regression checks passed.')
    p('Pre-scoring manifest SHA256: '+fingerprint)
    p('Project repository: https://github.com/Eillot1031/DASE7506-MP2')
    p('All asset sizes and full hashes are in experiments/optimization85/test-resources.json. Pre-scoring source and asset hashes are in experiments/optimization85/freeze.json. The earlier checkpoint and its six scoring records remain intact.')

    page('AI disclosure attribution and further work')
    sub('AI assistance disclosure')
    p('OpenAI Codex provided substantive assistance in reading the project requirements, implementing feature extraction and training code, running experiments, diagnosing validation overfitting, preparing reproducibility records, drafting this report and the README, and publishing code. The reported numerical scores come from execution of the unchanged independent course scorer. No score was fabricated or manually inserted into the prediction output.')
    p('The student author is responsible for reviewing and understanding the implementation, checking the analysis, and accurately disclosing the assistance and the adaptive testing history. No LLM generation or API call occurs during evaluation. No external training dataset, additional pretrained weights, question-ID feature or answer lookup is used. Evaluation prediction files are retained as evidence and are not read by inference.')
    sub('Attribution and provenance')
    p('The data, fixed representations, interface and official evaluation tools are reused from the course repository [1, 2]. Questions and references derive from GSM8K under the retained MIT license. Candidate traces and representations were prepared by the course with Qwen3-0.6B, revision c1899de289a04d12100db370d81485cdf75e47ca [3, 4]. No Qwen weights are distributed here. The original NOTICE and LICENSE remain in code/data/.')
    sub('Further work')
    p('An independent audit bank is the next defensible evaluation. Better grounding of entities and quantities and a held-out calibration procedure remain useful directions. The latest public score is 84.582%, exceeding the original 84% target but missing the subsequent 85% aspiration by two correct questions. Additional public-score tuning was stopped after the single frozen evaluation in this round; it would compound the adaptive testing limitation.')
    sub('References')
    p('[1] DASE7506 Mini Project 2 guide. https://github.com/xudongwu-0/DASE7506-MP2/blob/main/guide/GUIDE.md')
    p('[2] DASE7506 code and running instructions. https://github.com/xudongwu-0/DASE7506-MP2/blob/main/code/README.md')
    p('[3] Course data manifest, including generation revision and split hashes. code/data/manifest.json in the frozen repository version.')
    p('[4] GSM8K attribution and MIT license retained by the course. code/data/NOTICE.txt and code/data/LICENSE.txt in the frozen repository version.')
    p('[5] Zaheer et al. Deep Sets. NeurIPS 2017. https://arxiv.org/abs/1703.06114')
    p('[6] Foret et al. Sharpness-Aware Minimization for Efficiently Improving Generalization. ICLR 2021. https://arxiv.org/abs/2010.01412')
    # Remove inherited title rules and theme-font overrides in the default template.
    for style in d.styles:
        for border in list(style.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    for par in d.paragraphs:
        for border in list(par._p.xpath('.//w:pBdr')):border.getparent().remove(border)
        if par.style.name in ('Title','Subtitle','Heading 1','Heading 2'):
            for run in par.runs:
                run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0);run.font.underline=False
                fonts=run._r.get_or_add_rPr().rFonts
                for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
                    fonts.attrib.pop(qn('w:'+attr),None)
    target=ROOT/'report/Selecting_Reliable_LLM_Reasoning_Review.docx';d.save(target);print(target)

if __name__=='__main__':main()
